"""End-to-end API tests. A fake OpenAI-compatible server stands in for the real AI provider.

Run from the project folder:  python -m unittest discover -s tests -v
"""
from __future__ import annotations

import hashlib
import io
import json
import math
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

import httpx
from fastapi.testclient import TestClient

from backend.ai import AIClient
from backend.config import Settings
from backend.main import create_app
from backend.text import chunk_text, keyword_scores


def fake_vector(text: str, dims: int = 64) -> list[float]:
    """Deterministic bag-of-words vector, so similar texts get similar vectors."""
    vector = [0.0] * dims
    for word in text.lower().split():
        word = "".join(ch for ch in word if ch.isalnum())
        if len(word) > 3:
            vector[int(hashlib.md5(word.encode()).hexdigest(), 16) % dims] += 1.0
    return vector


class FakeProvider:
    def __init__(self, fail_chat: bool = False, fail_embed: bool = False):
        self.fail_chat, self.fail_embed = fail_chat, fail_embed
        self.calls: list[str] = []

    def __call__(self, request: httpx.Request) -> httpx.Response:
        path = request.url.path
        self.calls.append(path)
        if path.endswith("/models"):
            return httpx.Response(200, json={"data": [{"id": "fake-chat"}, {"id": "fake-embed"}]})
        body = json.loads(request.content)
        if path.endswith("/embeddings"):
            if self.fail_embed:
                return httpx.Response(500, json={"error": "boom"})
            return httpx.Response(200, json={"data": [
                {"index": i, "embedding": fake_vector(t)} for i, t in enumerate(body["input"])
            ]})
        if path.endswith("/chat/completions"):
            if self.fail_chat:
                raise httpx.ConnectError("down")
            user = body["messages"][1]["content"]
            assert "Retrieved company knowledge" in user
            return httpx.Response(200, json={"choices": [{"message": {"content": "Submit within 10 business days [S1]."}}]})
        return httpx.Response(404)


def make_client(tmp: str, provider: FakeProvider | None = None, **overrides) -> TestClient:
    base = Settings(data_dir=Path(tmp))
    if provider is not None:
        base = replace(base, provider="compatible", provider_label="Fake AI", base_url="http://fake/v1",
                       api_key="k", chat_model="fake-chat", embedding_model="fake-embed")
    settings = replace(base, **overrides)
    ai = AIClient(settings, transport=httpx.MockTransport(provider) if provider else None)
    return TestClient(create_app(settings, ai))


class TextTests(unittest.TestCase):
    def test_long_pdf_style_text_is_split(self):
        text = ("Employees must submit receipts within ten days of travel. " * 120).strip()  # one giant paragraph
        chunks = chunk_text(text)
        self.assertGreater(len(chunks), 4)
        self.assertTrue(all(len(c) <= 1200 for c in chunks))

    def test_keyword_scores_rank_the_matching_chunk_first(self):
        docs = ["Travel claims need receipts and a boarding pass.", "Git branches are created from main."]
        scores = keyword_scores("What do travel claims need?", docs)
        self.assertGreater(scores[0], scores[1])
        self.assertEqual(scores[1], 0.0)


class RetrievalOnlyTests(unittest.TestCase):
    """No AI configured: the app must still work and say so honestly."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.client = make_client(self.tmp.name).__enter__()

    def tearDown(self):
        self.client.__exit__(None, None, None)
        self.tmp.cleanup()

    def test_health_reports_retrieval_only(self):
        health = self.client.get("/api/health").json()
        self.assertFalse(health["api_configured"])
        self.assertEqual(health["mode"], "retrieval_only")
        self.assertEqual(health["indexed_chunks"], 4)  # the four sample documents

    def test_chat_returns_passages_with_sources(self):
        result = self.client.post("/api/chat", json={"question": "What documents are needed for travel reimbursement?"}).json()
        self.assertEqual(result["mode"], "retrieval_only")
        self.assertIn("AI answers are turned off", result["answer"])
        self.assertEqual(result["sources"][0]["document"], "Employee Travel & Expense Policy")

    def test_unrelated_question_gets_no_match(self):
        result = self.client.post("/api/chat", json={"question": "zebra xylophone quantum"}).json()
        self.assertEqual(result["mode"], "no_match")
        self.assertEqual(result["sources"], [])

    def test_secrets_are_refused(self):
        result = self.client.post("/api/chat", json={"question": "my api_key: sk-abcdefghijklmnopqrstuvwxyz"}).json()
        self.assertTrue(result["filtered"])

    def test_stats_are_real_not_hardcoded(self):
        stats = self.client.get("/api/stats").json()
        self.assertIsNone(stats["retrieval_hit_rate"])  # no questions yet
        self.client.post("/api/chat", json={"question": "travel reimbursement documents"})
        stats = self.client.get("/api/stats").json()
        self.assertEqual(stats["questions"], 1)
        self.assertNotIn("answer_acceptance", stats)

    def test_analytics_are_computed_from_real_questions(self):
        self.client.post("/api/chat", json={"question": "What documents are needed for travel reimbursement?"})
        self.client.post("/api/chat", json={"question": "travel expense claim deadline"})
        self.client.post("/api/chat", json={"question": "zebra xylophone quantum"})
        stats = self.client.get("/api/stats").json()
        self.assertEqual(len(stats["daily_questions"]), 14)
        today = stats["daily_questions"][-1]
        self.assertEqual((today["answered"], today["unanswered"]), (2, 1))
        top = stats["top_documents"][0]
        self.assertEqual(top["document"], "Employee Travel & Expense Policy")
        self.assertEqual(top["citations"], 2)
        self.assertIsNotNone(stats["median_response_ms"])

    def test_upload_markdown_then_delete(self):
        content = b"# Parking\n\nVisitors park in lot B and must register at reception."
        doc = self.client.post("/api/documents", files={"file": ("parking-rules.md", content, "text/markdown")}).json()
        self.assertEqual(doc["name"], "Parking Rules")
        answer = self.client.post("/api/chat", json={"question": "Where do visitors park?"}).json()
        self.assertEqual(answer["sources"][0]["document"], "Parking Rules")
        self.assertEqual(self.client.delete(f"/api/documents/{doc['id']}").status_code, 200)
        names = [d["name"] for d in self.client.get("/api/documents").json()]
        self.assertNotIn("Parking Rules", names)

    def test_upload_docx_with_table(self):
        from docx import Document

        document = Document()
        document.add_paragraph("Laptop refresh policy.")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Refresh cycle"
        table.rows[0].cells[1].text = "Every three years"
        buffer = io.BytesIO()
        document.save(buffer)
        doc = self.client.post("/api/documents", files={"file": ("laptops.docx", buffer.getvalue(), "application/octet-stream")}).json()
        self.assertEqual(doc["chunks"], 1)
        answer = self.client.post("/api/chat", json={"question": "laptop refresh cycle"}).json()
        self.assertIn("Every three years", answer["answer"])

    def test_upload_pdf(self):
        from pypdf import PdfWriter
        from pypdf.generic import NameObject, DecodedStreamObject, DictionaryObject

        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=200)
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 100 Td (Badge pickup happens at reception on day one.) Tj ET")
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        page[NameObject("/Contents")] = writer._add_object(stream)
        buffer = io.BytesIO()
        writer.write(buffer)
        doc = self.client.post("/api/documents", files={"file": ("badges.pdf", buffer.getvalue(), "application/pdf")})
        self.assertEqual(doc.status_code, 200, doc.text)
        answer = self.client.post("/api/chat", json={"question": "Where is badge pickup?"}).json()
        self.assertEqual(answer["sources"][0]["document"], "Badges")

    def test_rejects_unsupported_and_empty_files(self):
        self.assertEqual(self.client.post("/api/documents", files={"file": ("a.exe", b"MZ", "x")}).status_code, 400)
        self.assertEqual(self.client.post("/api/documents", files={"file": ("a.doc", b"x", "x")}).status_code, 400)
        self.assertEqual(self.client.post("/api/documents", files={"file": ("blank.md", b"   \n\n ", "x")}).status_code, 422)
        self.assertEqual(self.client.post("/api/documents", files={"file": ("bad.pdf", b"not a pdf", "x")}).status_code, 422)

    def test_history_is_private_per_visitor_in_demo_mode(self):
        with make_client(self.tmp.name + "-demo", demo_mode=True) as demo:
            demo.post("/api/chat", json={"question": "travel reimbursement"}, headers={"X-Client-Id": "visitor-aaaa"})
            mine = demo.get("/api/history", headers={"X-Client-Id": "visitor-aaaa"}).json()
            theirs = demo.get("/api/history", headers={"X-Client-Id": "visitor-bbbb"}).json()
        self.assertEqual(len(mine), 1)
        self.assertEqual(theirs, [])

    def test_demo_rate_limit_and_seed_protection(self):
        with make_client(self.tmp.name + "-limit", demo_mode=True, chat_per_hour=2) as demo:
            codes = [demo.post("/api/chat", json={"question": "travel policy"}).status_code for _ in range(3)]
            self.assertEqual(codes, [200, 200, 429])
            seed = demo.get("/api/documents").json()[0]
            self.assertEqual(demo.delete(f"/api/documents/{seed['id']}").status_code, 403)

    def test_demo_keeps_document_count_bounded(self):
        with make_client(self.tmp.name + "-prune", demo_mode=True, max_documents=5) as demo:
            for i in range(4):
                demo.post("/api/documents", files={"file": (f"note-{i}.md", f"Note number {i} about lunch.".encode(), "text/markdown")})
            names = [d["name"] for d in demo.get("/api/documents").json()]
        self.assertEqual(len(names), 5)  # 4 seeds + newest upload
        self.assertIn("Note 3", names)


class AIProviderTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.tmp.cleanup()

    def test_semantic_answer_with_citations(self):
        provider = FakeProvider()
        with make_client(self.tmp.name, provider) as client:
            health = client.get("/api/health").json()
            self.assertTrue(health["api_configured"])
            result = client.post("/api/chat", json={"question": "travel expense claim deadline business days"}).json()
        self.assertEqual(result["mode"], "ai")
        self.assertEqual(result["retrieval"], "semantic")
        self.assertIn("[S1]", result["answer"])
        self.assertEqual(result["sources"][0]["document"], "Employee Travel & Expense Policy")

    def test_vectors_are_cached_and_reembedded_when_model_changes(self):
        provider = FakeProvider()
        with make_client(self.tmp.name, provider) as client:
            client.post("/api/chat", json={"question": "travel claims"})
            client.post("/api/chat", json={"question": "release approvals"})
        # First question embeds all chunks + the query; the second embeds only the query.
        self.assertEqual(provider.calls.count("/v1/embeddings"), 3)

        provider2 = FakeProvider()
        with make_client(self.tmp.name, provider2, embedding_model="other-embed") as client:
            client.post("/api/chat", json={"question": "travel claims"})
        self.assertEqual(provider2.calls.count("/v1/embeddings"), 2)  # old vectors ignored, chunks re-embedded

    def test_falls_back_to_passages_when_chat_model_is_down(self):
        with make_client(self.tmp.name, FakeProvider(fail_chat=True)) as client:
            result = client.post("/api/chat", json={"question": "travel expense claim deadline"}).json()
        self.assertEqual(result["mode"], "retrieval_only")
        self.assertIn("unavailable right now", result["answer"])
        self.assertTrue(result["sources"])

    def test_upload_survives_embedding_outage(self):
        with make_client(self.tmp.name, FakeProvider(fail_embed=True)) as client:
            doc = client.post("/api/documents", files={"file": ("wifi.md", b"Guest wifi password rotates monthly.", "text/markdown")}).json()
            self.assertIn("warning", doc)
            result = client.post("/api/chat", json={"question": "guest wifi rotation"}).json()
        self.assertEqual(result["retrieval"], "keyword")
        self.assertEqual(result["sources"][0]["document"], "Wifi")


class SettingsTests(unittest.TestCase):
    def test_old_env_files_still_work(self):
        s = Settings.from_env({"OPENAI_API_KEY": "sk-test"})
        self.assertEqual(s.provider, "openai")
        self.assertTrue(s.ai_enabled)

    def test_openai_without_key_falls_back_to_retrieval_only(self):
        s = Settings.from_env({"AI_PROVIDER": "openai"})
        self.assertEqual(s.provider, "none")
        self.assertFalse(s.ai_enabled)

    def test_ollama_defaults(self):
        s = Settings.from_env({"AI_PROVIDER": "ollama"})
        self.assertEqual(s.base_url, "http://127.0.0.1:11434/v1")
        self.assertEqual(s.embedding_model, "nomic-embed-text")

    def test_ollama_latest_tag_counts_as_installed(self):
        settings = Settings.from_env({"AI_PROVIDER": "ollama"})

        def handler(request):
            return httpx.Response(200, json={"data": [{"id": "qwen2.5:7b"}, {"id": "nomic-embed-text:latest"}]})

        self.assertTrue(AIClient(settings, transport=httpx.MockTransport(handler)).ready())

    def test_demo_mode_turns_on_limits(self):
        s = Settings.from_env({"DEMO_MODE": "true"})
        self.assertEqual((s.max_upload_mb, s.chat_per_hour, s.max_documents), (5, 30, 25))


if __name__ == "__main__":
    unittest.main()
