"""Turning uploaded files into searchable text chunks, plus keyword scoring for retrieval without AI."""
from __future__ import annotations

import io
import logging
import math
import re
from collections import Counter
from pathlib import Path

logging.getLogger("pypdf").setLevel(logging.ERROR)  # malformed PDFs are reported to the user instead

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt"}


class ParseError(ValueError):
    pass


def extract_text(filename: str, raw: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension in {".md", ".markdown", ".txt"}:
        return raw.decode("utf-8", errors="ignore")
    if extension == ".pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise ParseError("This PDF is password-protected.")
            return "\n\n".join((page.extract_text() or "") for page in reader.pages)
        except ParseError:
            raise
        except Exception as exc:
            raise ParseError("This PDF could not be read. It may be damaged or scanned without text.") from exc
    if extension == ".docx":
        from docx import Document

        try:
            document = Document(io.BytesIO(raw))
        except Exception as exc:
            raise ParseError("This Word file could not be read. Save it as .docx and try again.") from exc
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:  # policies often keep key facts in tables
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n\n".join(parts)
    raise ParseError("Supported formats are PDF, DOCX, Markdown and plain text.")


_SENTENCE_END = re.compile(r"(?<=[.!?。！？])\s+")


def _split_long(paragraph: str, max_chars: int) -> list[str]:
    """Split one oversized paragraph at sentence boundaries, then hard-wrap anything still too long."""
    pieces: list[str] = []
    current = ""
    for sentence in _SENTENCE_END.split(paragraph):
        while len(sentence) > max_chars:  # a single run-on "sentence" (common in PDF text)
            cut = sentence.rfind(" ", 0, max_chars)
            cut = cut if cut > max_chars // 2 else max_chars
            pieces.append((current + " " + sentence[:cut]).strip() if current else sentence[:cut].strip())
            current = ""
            sentence = sentence[cut:].strip()
        if len(current) + len(sentence) + 1 <= max_chars:
            current = f"{current} {sentence}".strip()
        else:
            if current:
                pieces.append(current)
            current = sentence
    if current:
        pieces.append(current)
    return pieces


def chunk_text(text: str, max_chars: int = 1200, overlap: int = 180) -> list[str]:
    """Group paragraphs into chunks of at most `max_chars`, carrying a short overlap between chunks."""
    normalized = text.replace("\r\n", "\n").replace("\x00", "")
    normalized = re.sub(r"[ \t]+", " ", normalized)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized).strip()
    if not normalized:
        return []
    paragraphs: list[str] = []
    for part in re.split(r"\n\s*\n", normalized):
        part = part.strip()
        if not part:
            continue
        paragraphs.extend(_split_long(part, max_chars - overlap) if len(part) > max_chars - overlap else [part])

    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(current) + len(paragraph) + 2 <= max_chars:
            current = f"{current}\n\n{paragraph}".strip()
            continue
        if current:
            chunks.append(current)
        tail = current[-overlap:] if current else ""
        if tail and " " in tail:
            tail = tail[tail.index(" ") + 1 :]  # start the overlap on a word boundary
        current = f"{tail}\n\n{paragraph}".strip() if len(tail) + len(paragraph) + 2 <= max_chars else paragraph
    if current:
        chunks.append(current)
    return chunks


_STOPWORDS = set(
    """a an and are as at be but by can do does for from has have how i if in into is it its may of on or
    our should so than that the their them then there these they this to was we what when where which who
    why will with you your need needs""".split()
)
_TOKEN = re.compile(r"[a-z0-9]+|[一-鿿]")


def tokenize(text: str) -> list[str]:
    tokens = []
    for token in _TOKEN.findall(text.lower()):
        if len(token) == 1 and token.isascii():
            continue
        if token in _STOPWORDS:
            continue
        if len(token) > 4 and token.endswith("s") and not token.endswith("ss"):
            token = token[:-1]  # crude plural folding: "documents" -> "document"
        tokens.append(token)
    return tokens


def keyword_scores(query: str, documents: list[str]) -> list[float]:
    """BM25 scores squashed into 0..1 so they read like the cosine scores used with AI embeddings."""
    query_terms = set(tokenize(query))
    if not query_terms or not documents:
        return [0.0] * len(documents)
    tokenized = [tokenize(doc) for doc in documents]
    avg_len = sum(len(t) for t in tokenized) / len(tokenized) or 1.0
    doc_freq = Counter(term for tokens in tokenized for term in set(tokens))
    n = len(documents)
    k1, b = 1.4, 0.75
    raw_scores = []
    for tokens in tokenized:
        counts = Counter(tokens)
        score = 0.0
        for term in query_terms:
            tf = counts.get(term, 0)
            if not tf:
                continue
            idf = math.log(1 + (n - doc_freq[term] + 0.5) / (doc_freq[term] + 0.5))
            score += idf * tf * (k1 + 1) / (tf + k1 * (1 - b + b * len(tokens) / avg_len))
        raw_scores.append(score)
    coverage = [len(query_terms & set(tokens)) / len(query_terms) for tokens in tokenized]
    top = max(raw_scores) or 1.0
    # Blend relative BM25 strength with how many query words were found at all.
    return [round(0.5 * (s / top) + 0.5 * c, 4) if s else 0.0 for s, c in zip(raw_scores, coverage)]
